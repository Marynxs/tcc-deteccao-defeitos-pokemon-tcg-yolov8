"""Numera as cartas novas, atualiza a planilha e gera a pasta de upload do Roboflow.

A entrega nova usa outro layout de pasta e outro formato de arquivo que a
primeira (front_image.webp em vez de frente.png), e veio com cartas de outros
jogos misturadas, por isso este script existe separado de organize_cards.py.

Reconstroi as linhas novas da planilha do zero a cada execucao, para que o ID
sequencial e as observacoes sempre correspondam ao que foi de fato exportado.
"""

import hashlib
import json
import shutil
from copy import copy
from pathlib import Path

import openpyxl
from PIL import Image

PLANILHA = Path("Cartas TCG.xlsx")
ABA = "Mapeamento"
ENTRADA = Path("grading_reports_photos")
SAIDA = Path("Dataset_YOLO/upload_roboflow")
SAIDA_FORA = SAIDA / "_fora_do_escopo"
LINKS = Path("docs/links_registro.json")

BASE_URL = "https://sites.google.com/view/gradingcapy/"
PRIMEIRA_LINHA_NOVA = 84  # a linha 83 e a carta_082, ultima do lote ja anotado
LADOS = {"frente": "front_image.webp", "verso": "back_image.webp"}

# Cartas que nao sao Pokemon TCG, identificadas pelo verso. Ficam separadas em
# vez de apagadas: o dominio do trabalho e carta Pokemon, mas elas servem como
# conjunto externo se algum dia interessar medir generalizacao entre jogos.
# A Ancient Mew (C001115) chegou a entrar nesta lista por engano: o verso
# dourado e exclusivo dessa promo e nao lembra o verso comum da Pokemon.
FORA_DO_ESCOPO = {
    "C001050": "Magic: The Gathering", "C001058": "Magic: The Gathering",
    "C001069": "Magic: The Gathering", "C001072": "Magic: The Gathering",
    "C001074": "Magic: The Gathering", "C001153": "Magic: The Gathering",
    "C001154": "Magic: The Gathering", "C001155": "Magic: The Gathering",
    "C001156": "Magic: The Gathering", "C001157": "Magic: The Gathering",
    "C001158": "Magic: The Gathering", "C001159": "Magic: The Gathering",
    "C001160": "Magic: The Gathering", "C001161": "Magic: The Gathering",
    "C001162": "Magic: The Gathering", "C001198": "Magic: The Gathering",
    "C001245": "Magic: The Gathering",
    "C001195": "card Marvel", "C001196": "card Marvel", "C001197": "card Marvel",
    "C001211": "One Piece Card Game", "C001225": "One Piece Card Game",
    "C001224": "outro TCG",
    "C001257": "baralho antigo, nao e TCG",
}

# Pasta que veio junto mas nao corresponde a uma carta: e copia byte a byte de
# outra da propria entrega.
NAO_E_CARTA = {"Teste": "Copia byte a byte de C001126"}


def pastas_por_codigo():
    """Indexa as pastas pelo codigo limpo.

    Uma das pastas veio como "report_ C001222", com espaco antes do codigo,
    entao montar o caminho concatenando o codigo perderia essa carta."""
    return {
        d.name.replace("report_", "").strip(): d
        for d in ENTRADA.iterdir()
        if d.is_dir()
    }


def imagens(pasta):
    if pasta is None:
        return {}
    return {lado: pasta / arq for lado, arq in LADOS.items() if (pasta / arq).is_file()}


def duplicatas(codigos, pastas):
    """Codigos cujas imagens repetem as de um codigo anterior.

    Duas cartas com a mesma foto cairiam em folds diferentes e vazariam
    informacao entre treino e validacao, entao a segunda ocorrencia fica de
    fora ate a empresa esclarecer."""
    visto, repetidos = {}, {}
    for codigo in codigos:
        for caminho in imagens(pastas.get(codigo)).values():
            digest = hashlib.sha256(caminho.read_bytes()).hexdigest()
            if digest in visto and visto[digest] != codigo:
                repetidos[codigo] = visto[digest]
            visto.setdefault(digest, codigo)
    return repetidos


def links_quebrados():
    if not LINKS.is_file():
        return {}
    dados = json.loads(LINKS.read_text(encoding="utf-8"))
    return {d["codigo"]: d["status"] for d in dados if d["status"] != 200}


def classificar(codigo, pasta, repetidos, quebrados):
    """Devolve (entra_no_dataset, lista de observacoes)."""
    notas, entra = [], True
    presentes = imagens(pasta)

    if codigo in NAO_E_CARTA:
        return False, [f"{NAO_E_CARTA[codigo]}. Confirmar com a empresa."]
    if not presentes:
        entra = False
        notas.append("Pasta sem imagens")
    elif len(presentes) < 2:
        entra = False
        notas.append(f"So um lado disponivel: {sorted(presentes)}")
    if codigo in repetidos:
        entra = False
        notas.append(f"Imagens identicas as de {repetidos[codigo]}")
    if codigo in FORA_DO_ESCOPO:
        entra = False
        notas.append(f"Nao e carta Pokemon ({FORA_DO_ESCOPO[codigo]})")
    if pasta is not None and pasta.name != f"report_{codigo}":
        notas.append(f"Nome da pasta irregular: {pasta.name!r}")
    if codigo in quebrados:
        notas.append(f"Link do registro fora do ar ({quebrados[codigo]})")
    return entra, notas


def copiar_estilo(ws, destino, coluna, modelo_linha):
    modelo = ws.cell(modelo_linha, coluna)
    alvo = ws.cell(destino, coluna)
    alvo.font = copy(modelo.font)
    alvo.alignment = copy(modelo.alignment)
    alvo.border = copy(modelo.border)
    alvo.number_format = modelo.number_format
    return alvo


def main():
    wb = openpyxl.load_workbook(PLANILHA)
    ws = wb[ABA]

    ja_na_planilha = {
        ws.cell(r, 2).value.strip(): ws.cell(r, 1).value
        for r in range(2, PRIMEIRA_LINHA_NOVA)
        if ws.cell(r, 2).value
    }
    pastas = pastas_por_codigo()
    novos = sorted(c for c in pastas if c not in ja_na_planilha)
    repetidos = duplicatas(novos, pastas)
    quebrados = links_quebrados()

    for r in range(PRIMEIRA_LINHA_NOVA, ws.max_row + 1):
        for c in range(1, 8):
            ws.cell(r, c).value = None

    if SAIDA.exists():
        shutil.rmtree(SAIDA)
    SAIDA.mkdir(parents=True)
    SAIDA_FORA.mkdir()

    proximo = max(int(v.split("_")[1]) for v in ja_na_planilha.values()) + 1
    linha = PRIMEIRA_LINHA_NOVA - 1
    dentro, foraram = [], []

    for codigo in novos:
        linha += 1
        pasta = pastas.get(codigo)
        presentes = imagens(pasta)
        entra, notas = classificar(codigo, pasta, repetidos, quebrados)

        # So quem entra no dataset consome um carta_NNN, para que a sequencia
        # fique sem buraco. O que fica de fora continua identificado pelo
        # codigo original, que e a identidade real da carta.
        ident = f"carta_{proximo:03d}" if entra else None
        if entra:
            proximo += 1
            dentro.append((ident, codigo))
        else:
            foraram.append((codigo, "; ".join(notas)))

        for c in range(1, 8):
            copiar_estilo(ws, linha, c, PRIMEIRA_LINHA_NOVA - 1).fill = openpyxl.styles.PatternFill(fill_type=None)
        ws.cell(linha, 1).value = ident
        ws.cell(linha, 2).value = codigo
        ws.cell(linha, 3).value = f'=HYPERLINK("{BASE_URL}{codigo}")'
        ws.cell(linha, 4).value = None
        ws.cell(linha, 5).value = "OK" if "frente" in presentes else None
        ws.cell(linha, 6).value = "OK" if "verso" in presentes else None
        ws.cell(linha, 7).value = "; ".join(notas) or None

        destino = SAIDA if entra else SAIDA_FORA
        nome = ident or codigo
        for lado, origem in presentes.items():
            # A conversao para PNG nao recupera nada do que o webp com perdas ja
            # descartou; serve para nao acrescentar perda nova e para manter a
            # mesma extensao do lote que ja esta no Roboflow.
            Image.open(origem).convert("RGB").save(destino / f"{nome}_{lado}.png")

    wb.save(PLANILHA)

    print(f"pastas lidas:        {len(pastas)}")
    print(f"ja anotadas:         {len(ja_na_planilha)}")
    print(f"novas no dataset:    {len(dentro)}  ({dentro[0][0]} a {dentro[-1][0]})")
    print(f"total do dataset:    {len(ja_na_planilha) + len(dentro)} cartas")
    print(f"\nfora do dataset: {len(foraram)}")
    for codigo, motivo in foraram:
        print(f"  {codigo}  {motivo}")


if __name__ == "__main__":
    main()
