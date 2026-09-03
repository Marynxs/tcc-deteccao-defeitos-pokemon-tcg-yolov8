import os
import shutil
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ─── EDITE APENAS ESTAS LINHAS ─────────────────────────────────────────────
ROOT_DIR  = "Dataset"        # folder que contém C001013, C001014...
OUTPUT_DIR = "Dataset_TCC2"
BASE_URL    = "https://sites.google.com/view/gradingcapy/"
# ───────────────────────────────────────────────────────────────────────────

# Candidatos aceitos para frente e verso, em ordem de prioridade
CANDIDATES = {
    "frente": ["frente.png", "1.png"],
    "verso":  ["verso.png",  "2.png"],
}

def find_image(folder: Path, side: str) -> Path | None:
    """Retorna o primeiro candidate encontrado para frente ou verso."""
    for name in CANDIDATES[side]:
        candidate = folder / name
        if candidate.exists():
            return candidate
    return None

def clean_output(output_folder: Path, xlsx_path: Path):
    """Remove a folder de saída e a planilha anteriores para garantir execução limpa."""
    if output_folder.exists():
        shutil.rmtree(output_folder)
        print(f"🗑️   Pasta de saída anterior removida: {output_folder}")

    if xlsx_path.exists():
        xlsx_path.unlink()
        print(f"🗑️   Planilha anterior removida: {xlsx_path}")

def organize():
    root_folder  = Path(ROOT_DIR)
    output_folder = Path(OUTPUT_DIR)
    xlsx_path   = output_folder.parent / "mapeamento_cartas.xlsx"

    # ── Limpa execução anterior antes de qualquer coisa ──────────────────────
    clean_output(output_folder, xlsx_path)
    output_folder.mkdir(parents=True, exist_ok=True)

    # Pega todas as subfolders (C001013, C001014 ...) em ordem alfabética
    subfolders = sorted([p for p in root_folder.iterdir() if p.is_dir()])

    copied = 0
    errors    = []
    mapping = []

    for idx, folder in enumerate(subfolders, start=1):
        number = f"{idx:03d}"       # 001, 002, 003 ...
        code = folder.name         # C001013
        link   = BASE_URL + code  # URL completa do registro

        line = {
            "carta_id":      f"carta_{number}",
            "codigo_orig":   code,
            "link_registro": link,
            "frente_ok":     "",
            "verso_ok":      "",
        }

        for side in ["frente", "verso"]:
            source  = find_image(folder, side)
            target = output_folder / f"carta_{number}_{side}.png"

            if source is None:
                msg = f"AUSENTE: {code} — nenhum candidate encontrado para '{side}' {CANDIDATES[side]}"
                errors.append(msg)
                line[f"{side}_ok"] = "AUSENTE"
                continue

            shutil.copy2(source, target)
            line[f"{side}_ok"] = "OK"
            copied += 1
            print(f"✓  {code}/{source.name:<20} →  carta_{number}_{side}.png")

        mapping.append(line)

    # ── Salva planilha .xlsx (sem risco de delimitador errado) ──────────────
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mapeamento"

    fields = ["carta_id", "codigo_orig", "link_registro", "frente_ok", "verso_ok"]
    headers = ["ID da Carta", "Código Original", "Link do Registro", "Frente", "Verso"]

    # Cabeçalho estilizado
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1D9E75", end_color="1D9E75", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    # Linhas de dados
    for line in mapping:
        ws.append([line[c] for c in fields])

    # Fonte Arial em todas as células de dados
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = Font(name="Arial")

    # Larguras de coluna razoáveis
    widths = {"A": 14, "B": 16, "C": 48, "D": 12, "E": 12}
    for col, largura in widths.items():
        ws.column_dimensions[col].width = largura

    wb.save(xlsx_path)

    # ── Resumo ───────────────────────────────────────────────────────────────
    print(f"\n{'='*55}")
    print(f"✅  {copied} images copied")
    print(f"📄  Planilha salva em: {xlsx_path}")

    if errors:
        print(f"\n⚠️   {len(errors)} problema(s) encontrado(s):")
        for e in errors:
            print(f"    {e}")
    else:
        print("✅  Nenhum problema encontrado")

if __name__ == "__main__":
    organize()